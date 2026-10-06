# PLOS ONE layout: Times New Roman 12 pt, double spacing, continuous line numbers, page numbers,
# single-column A4, tables styled as in the journal-neutral version (single-spaced cells).
import sys, os, re, copy
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import style_docx

def make_reference(src, out):
    d = Document(src)
    st = d.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Abstract", "Bibliography"):
        try: s = st[name]
        except KeyError: continue
        s.font.name = "Times New Roman"; s.font.size = Pt(12)
        s.paragraph_format.line_spacing = 2.0; s.paragraph_format.space_after = Pt(0); s.paragraph_format.space_before = Pt(0)
    for lvl, size in ((1, 14), (2, 12), (3, 12)):
        try: s = st[f"Heading {lvl}"]
        except KeyError: continue
        s.font.name = "Times New Roman"; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.italic = (lvl == 3)
        s.paragraph_format.line_spacing = 2.0; s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(0)
    for sec in d.sections:
        sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, m, Cm(2.5))
    d.save(out)

def add_page_number(section):
    footer = section.footer; p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = 1
    p._p.get_or_add_pPr().append(OxmlElement("w:suppressLineNumbers"))
    r = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)

def finish(path, out):
    tmp = out + ".tmp.docx"
    style_docx.style_tables(path, tmp)
    d = Document(tmp)
    for sec in d.sections:
        sp = sec._sectPr
        ln = sp.find(qn("w:lnNumType"))
        if ln is None: ln = OxmlElement("w:lnNumType"); sp.append(ln)
        ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "283")
        cols = sp.find(qn("w:cols"))
        if cols is not None: cols.set(qn("w:num"), "1")
        add_page_number(sec)
    # body text double-spaced 12 pt; table cells keep single spacing; table titles bold 12 pt; legends below tables 11 pt
    for p in d.paragraphs:
        pf = p.paragraph_format
        pf.line_spacing = 2.0
        if re.match(r"^Table \d+\. ", p.text):
            for r in p.runs: r.font.size = Pt(12); r.font.bold = True
    body = d.element.body
    for tb in d.tables:  # legend paragraph directly after a table: 11 pt
        nx = tb._tbl.getnext()
        if nx is not None and nx.tag == qn("w:p"):
            for rr in nx.iter(qn("w:r")):
                rpr = rr.find(qn("w:rPr"))
                if rpr is None: rpr = OxmlElement("w:rPr"); rr.insert(0, rpr)
                sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22"); rpr.append(sz)
    # suppress line numbers inside tables (PLOS counts body lines)
    for tb in d.tables:
        for p in tb._tbl.iter(qn("w:p")):
            ppr = p.find(qn("w:pPr"))
            if ppr is None: ppr = OxmlElement("w:pPr"); p.insert(0, ppr)
            if ppr.find(qn("w:suppressLineNumbers")) is None: ppr.append(OxmlElement("w:suppressLineNumbers"))
    d.save(out); os.remove(tmp)

if __name__ == "__main__":
    if sys.argv[1] == "ref": make_reference(sys.argv[2], sys.argv[3])
    else: finish(sys.argv[2], sys.argv[3])
