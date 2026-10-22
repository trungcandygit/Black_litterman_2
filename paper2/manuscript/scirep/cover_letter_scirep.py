"""Cover letter for Scientific Reports (items the journal asks for; unknown facts are left as bracketed placeholders)."""
import os
from docx import Document
from docx.shared import Pt, Cm
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "final", "SciRep", "Cover_letter.docx")
d = Document()
s = d.sections[0]; s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
P = lambda t: d.add_paragraph(t)
P("The Editors\nScientific Reports")
P("Dear Editors,")
P("Please consider our manuscript, “Do price limits delay price discovery? Overnight gaps after limit closes in Vietnam”, for publication as an Article in Scientific Reports.")
P("Corresponding author: Nguyen Van Trung, Academy of Policy and Development, Nam An Khanh Urban Area, Hoai Duc District, Hanoi, Vietnam; e-mail: kontrungcany@gmail.com.")
P("Why Scientific Reports. Price limits are used by exchanges worldwide, yet the data show little about how prices adjust after a stock closes at its limit. We split the next-day return after limit closes on the Ho Chi Minh Stock Exchange into an overnight gap and an intraday return, compare limit closes with stocks that moved almost as far, and test the result against 22 other price- and volume-based signals under false discovery rate control. We report the results as they are, including the floor-close result that depends on the benchmark and the explanations that our data cannot separate. The study is a technically sound empirical analysis based on public data and openly available code, which matches the journal's criteria of scientific validity rather than perceived impact, and it will interest readers in economics, finance and market design.")
P("The manuscript is original, is not under consideration elsewhere, and all authors have approved it. Generative AI use is documented in the Methods section. The data availability statement, author contributions and competing interests statement are included in the manuscript.")
P("Suggested reviewers: [authors to supply names, affiliations and e-mail addresses].")
P("Excluded referees: [authors to supply names, or state none].")
P("Prior discussion with an Editorial Board Member: [authors to confirm; state none if there was no discussion].")
P("Sincerely,\nNguyen Van Trung, on behalf of all authors")
d.save(OUT)
