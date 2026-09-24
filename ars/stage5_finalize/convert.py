"""Stage 5 conversion: Markdown -> DOCX (pandoc, native Word equations) and -> LaTeX (elsarticle) -> PDF (xelatex).

Run after build_stage5.py. Tectonic's bundle host is unreachable from this environment, so the
PDF is compiled with XeLaTeX from TeX Live (formatter_agent: "tectonic or xelatex").
"""
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
WORK = HERE / "work"
OUT = HERE.parents[1] / "submission"


def run(cmd, cwd=WORK):
    r = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout[-3000:], r.stderr[-3000:])
        raise SystemExit(f"failed: {cmd}")
    return r


# ---------------------------------------------------------------- reference.docx
ref = WORK / "reference.docx"
subprocess.run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"], check=True)
doc = Document(ref)
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
    # page number, top right
    p = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
    p.alignment = 2
    r = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); r._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; r._r.append(el)
for st in doc.styles:
    try:
        f = st.font
    except AttributeError:
        continue
    f.name = "Times New Roman"
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
for name, size in (("Normal", 12), ("Body Text", 12), ("First Paragraph", 12), ("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12), ("Compact", 10), ("Table", 10), ("Image Caption", 11)):
    try:
        st = doc.styles[name]
    except KeyError:
        continue
    st.font.size = Pt(size)
    if name.startswith("Heading") or name == "Title":
        st.font.bold = True
        st.font.color.rgb = None
    if name in ("Normal", "Body Text", "First Paragraph"):
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        st.paragraph_format.space_after = Pt(0)
doc.save(ref)

# ---------------------------------------------------------------- DOCX
for stem in ("manuscript_anonymized", "title_page", "highlights", "cover_letter", "declaration_of_competing_interest"):
    run(["pandoc", f"{stem}.md", "-f", "markdown+tex_math_dollars+superscript", "-o", f"{stem}.docx",
         "--reference-doc", "reference.docx", "--resource-path", "."])

# ---------------------------------------------------------------- LaTeX (elsarticle, review mode, anonymized)
md = (WORK / "manuscript_anonymized.md").read_text()
title = re.search(r"^# (.+)$", md, re.M).group(1)
abstract = re.search(r"## Abstract\n+(.+?)\n", md, re.S).group(1).strip()
keywords = [k.strip() for k in re.search(r"\*\*Keywords\*\*: (.+)", md).group(1).split(";")]
jel = [k.strip() for k in re.search(r"\*\*JEL classification\*\*: (.+)", md).group(1).split(";")]
body_md = md.split("**JEL classification**")[1].split("\n", 1)[1]
body_md = re.sub(r"^## ", "# ", body_md, flags=re.M)
body_md = re.sub(r"^### ", "## ", body_md, flags=re.M)
(WORK / "body.md").write_text(body_md)


def tex_inline(text):
    r = subprocess.run(["pandoc", "-f", "markdown", "-t", "latex"], input=text, capture_output=True, text=True, check=True)
    return r.stdout.strip()


title = tex_inline(title)
abstract = tex_inline(abstract)
keywords = [tex_inline(k) for k in keywords]
run(["pandoc", "body.md", "-f", "markdown+tex_math_dollars", "-t", "latex", "-o", "body.tex", "--top-level-division=section"])
body_tex = (WORK / "body.tex").read_text()
# unnumbered headings are already numbered in the text ("1. Introduction") -> use starred sections
body_tex = re.sub(r"\\(section|subsection)\{", r"\\\1*{", body_tex)
body_tex = re.sub(r"\\label\{[^}]*\}", "", body_tex)
body_tex = body_tex.replace(r"\includegraphics[width=1\textwidth,height=\textheight]", r"\noindent\includegraphics[width=0.95\textwidth]")
# keep figure title, image and note together; keep table titles with their tables
body_tex = re.sub(r"(\\textbf\{Figure \d\..*?\})\n\n(\\noindent\\includegraphics\[[^]]*\]\{[^}]*\})\n\n(\\emph\{Note\}:.*?)\n\n",
                  lambda m: "\\begin{figure}[p]\n\\noindent " + m.group(1) + "\\par\\medskip\n" + m.group(2) + "\\par\\medskip\n{\\small " + m.group(3) + "}\n\\end{figure}\n\n",
                  body_tex, flags=re.S)
body_tex = re.sub(r"\n(\\textbf\{Table (?:\d+|A\.\d)\.)", r"\n\\needspace{14\\baselineskip}\1", body_tex)
# table column widths (formatter_agent rule): widen the label column, share the rest
def colspec(m):
    inner = m.group(0)[len(r"\begin{longtable}[]{@{}"):-len("@{}}")]
    n = inner.count("p{") or len(re.findall(r"[lrc]", inner))
    first = 0.34 if n <= 4 else (0.26 if n <= 6 else 0.20)
    rest = (1 - first) / (n - 1)
    gaps = 2 * (n - 1)
    cols = [first] + [rest] * (n - 1)
    spec = "".join(r">{\raggedright\arraybackslash}p{(\linewidth - %d\tabcolsep) * \real{%.4f}}" % (gaps, c) for c in cols)
    return r"\begin{longtable}[]{@{}" + spec + "@{}}"
body_tex = re.sub(r"\\begin\{longtable\}\[\]\{@\{\}.*?@\{\}\}", colspec, body_tex, flags=re.S)
# break long source paths and URLs (xurl)
body_tex = re.sub(r"(?<![{/])\b(output/[A-Za-z0-9_\\/.]+?\.csv)", lambda m: r"\url{" + m.group(1).replace("\\_", "_") + "}", body_tex)
body_tex = re.sub(r"(?<![{])(https?://[^\s}]+[^\s}.,;)])", lambda m: r"\url{" + m.group(1).replace("\\_", "_").replace("\\%", "%").replace("\\&", "&").replace("\\#", "#") + "}", body_tex)
tex = r"""\documentclass[review,authoryear,12pt]{elsarticle}
\usepackage{fontspec}
\setmainfont{TeX Gyre Termes}
\usepackage{amsmath,amssymb,graphicx,booktabs,longtable,array,calc,etoolbox}
\usepackage[hidelinks]{hyperref}
\usepackage{xurl}
\usepackage{setspace}
\usepackage{needspace}
\providecommand{\tightlist}{\setlength{\itemsep}{0pt}\setlength{\parskip}{0pt}}
\providecommand{\pandocbounded}[1]{#1}
\AtBeginEnvironment{longtable}{\footnotesize\setstretch{1.0}}
\setlength{\LTleft}{0pt}
\setlength{\emergencystretch}{3em}
\journal{Finance Research Open}
\begin{document}
\begin{frontmatter}
\title{""" + title + r"""}
\begin{abstract}
""" + abstract + r"""
\end{abstract}
\begin{keyword}
""" + r" \sep ".join(keywords) + r"""
\JEL """ + r" \sep ".join(jel) + r"""
\end{keyword}
\end{frontmatter}
\doublespacing
""" + body_tex + r"""
\end{document}
"""
tex = tex.replace("%", r"\%").replace(r"\\%", r"\%") if False else tex
(WORK / "manuscript_anonymized.tex").write_text(tex)
for _ in range(2):
    run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", "manuscript_anonymized.tex"])
log = (WORK / "manuscript_anonymized.log").read_text(errors="ignore")
overfull = re.findall(r"Overfull \\hbox \((\d+\.\d+)pt too wide\)", log)
print("overfull boxes > 10pt:", [x for x in overfull if float(x) > 10])

# ---------------------------------------------------------------- copy deliverables
for f in ("manuscript_anonymized.docx", "manuscript_anonymized.tex", "manuscript_anonymized.pdf", "title_page.docx",
          "highlights.docx", "cover_letter.docx", "declaration_of_competing_interest.docx", "manuscript_anonymized.md"):
    shutil.copy(WORK / f, OUT / f)
print("done")
