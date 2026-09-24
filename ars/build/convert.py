"""Stage 5 conversion: Markdown -> DOCX (pandoc, native Word equations) and -> LaTeX (elsarticle) -> PDF (xelatex).

Run after build_stage5.py. Tectonic's bundle host is unreachable from this environment, so the
PDF is compiled with XeLaTeX from TeX Live (formatter_agent: "tectonic or xelatex").
"""
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from docx_schema_fix import fix as schema_fix

HERE = Path(__file__).resolve().parent
WORK = HERE / "work"
OUT = HERE.parents[1] / "submission"


def run(cmd, cwd=WORK):
    r = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout[-3000:], r.stderr[-3000:])
        raise SystemExit(f"failed: {cmd}")
    return r


# ---------------------------------------------------------------- reference.docx
ref = WORK / "reference.docx"
subprocess.run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"], check=True)
doc = Document(ref)
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
    # page number, top right
    p = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
    p.alignment = 2
    r = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); r._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; r._r.append(el)
for st in doc.styles:
    try:
        f = st.font
    except AttributeError:
        continue
    f.name = "Times New Roman"
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
for name, size in (("Normal", 12), ("Body Text", 12), ("First Paragraph", 12), ("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12), ("Compact", 10), ("Table", 10), ("Image Caption", 11)):
    try:
        st = doc.styles[name]
    except KeyError:
        continue
    st.font.size = Pt(size)
    if name.startswith("Heading") or name == "Title":
        st.font.bold = True
    if name in ("Normal", "Body Text", "First Paragraph"):
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
# journal style: all text black (pandoc's default styles use blue theme colours for titles,
# headings and hyperlinks); titles centred, headings left, body justified
for st in doc.styles:
    rpr = st.element.find(qn("w:rPr"))
    if rpr is None:
        continue
    for c in rpr.findall(qn("w:color")):
        rpr.remove(c)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "000000"); rpr.append(col)
for name in ("Title", "Subtitle"):
    try:
        doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except KeyError:
        pass
for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
    try:
        doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    except KeyError:
        pass
try:
    doc.styles["Hyperlink"].font.underline = False
except KeyError:
    pass
for name in ("Compact", "Table"):
    try:
        pf = doc.styles[name].paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(1); pf.space_after = Pt(1)
    except KeyError:
        pass
doc.save(ref)


def finish_docx(path, line_numbers=False, justify=True):
    """Post-process a pandoc DOCX: APA hanging indent for the reference list (left-aligned,
    so long DOIs and URLs do not stretch lines), table text left-aligned."""
    d = Document(path)
    if line_numbers:  # continuous line numbers for review copies (Elsevier recommendation)
        for sec in d.sections:
            for hp in sec.header.paragraphs:  # page-number header is not counted as a line
                ppr = hp._p.get_or_add_pPr()
                if ppr.find(qn("w:suppressLineNumbers")) is None:
                    ppr.append(OxmlElement("w:suppressLineNumbers"))
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "360")
            pg = sec._sectPr.find(qn("w:pgMar"))
            pg.addnext(ln) if pg is not None else sec._sectPr.append(ln)
    for para in d.paragraphs:  # keep table/figure titles and panel labels with what follows
        tx = para.text.strip()
        if re.match(r"^(Table (\d+|A\.\d)\.|Figure \d\.|Panel [A-C]\.)", tx):
            para.paragraph_format.keep_with_next = True
    if not justify:  # letters and short forms: ragged right
        for para in d.paragraphs:
            if para.paragraph_format.alignment in (None, WD_ALIGN_PARAGRAPH.JUSTIFY):
                para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    in_refs = False
    for para in d.paragraphs:
        if para.style.name.startswith("Heading"):
            in_refs = para.text.strip() == "References"
            continue
        if in_refs and para.text.strip():
            pf = para.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.left_indent = Cm(1.27)
            pf.first_line_indent = Cm(-1.27)
    for para in d.paragraphs:
        if para.style.name.startswith("Heading") or para.style.name in ("Title", "Subtitle"):
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.paragraph_format.alignment in (None, WD_ALIGN_PARAGRAPH.JUSTIFY):
                        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # column widths in proportion to the longest word and the typical cell length, so the
        # label column is not squeezed (pandoc gives every column the same width)
        ncol = len(t.columns)
        cw = 105  # approx. twips per character at 10 pt
        mins, wants = [], []
        for j in range(ncol):
            texts = [row.cells[j].text for row in t.rows if j < len(row.cells)]
            longest_word = max((len(w) for x in texts for w in x.split()), default=4)
            longest_cell = max((len(x) for x in texts), default=4)
            cap = 26 if j == 0 else 18  # label column may wrap between words; estimate cells stay on one line
            mins.append((longest_word + 2) * cw)
            wants.append((max(min(longest_cell, cap), longest_word) + 2) * cw)
        total = 9000  # text width in twips (A4, 2.54 cm margins)
        if sum(wants) <= total:
            widths = [int(w * total / sum(wants)) for w in wants]
        else:
            spare = max(total - sum(mins), 0)
            extra = [w - m for w, m in zip(wants, mins)]
            widths = [int(m + (e * spare / sum(extra) if sum(extra) else 0)) for m, e in zip(mins, extra)]
        tbl = t._tbl
        tblPr = tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(total))
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
        for gc, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
                tcW.set(qn("w:type"), "dxa"); tcW.set(qn("w:w"), str(w))
    d.save(path)
    schema_fix(path)

# ---------------------------------------------------------------- DOCX
for stem in ("manuscript_anonymized", "title_page", "highlights", "cover_letter", "declaration_of_competing_interest", "manuscript_with_author_details"):
    run(["pandoc", f"{stem}.md", "-f", "markdown+tex_math_dollars+superscript", "-o", f"{stem}.docx",
         "--reference-doc", "reference.docx", "--resource-path", "."])
    finish_docx(WORK / f"{stem}.docx", line_numbers=stem.startswith("manuscript"), justify=stem.startswith("manuscript"))

# ---------------------------------------------------------------- LaTeX (elsarticle, review mode, anonymized)
md = (WORK / "manuscript_anonymized.md").read_text()
title = re.search(r"^# (.+)$", md, re.M).group(1)
abstract = re.search(r"## Abstract\n+(.+?)\n", md, re.S).group(1).strip()
keywords = [k.strip() for k in re.search(r"\*\*Keywords\*\*: (.+)", md).group(1).split(";")]
jel = [k.strip() for k in re.search(r"\*\*JEL classification\*\*: (.+)", md).group(1).split(";")]
body_md = md.split("**JEL classification**")[1].split("\n", 1)[1]
body_md = re.sub(r"^## ", "# ", body_md, flags=re.M)
body_md = re.sub(r"^### ", "## ", body_md, flags=re.M)
(WORK / "body.md").write_text(body_md)


def tex_inline(text):
    r = subprocess.run(["pandoc", "-f", "markdown", "-t", "latex"], input=text, capture_output=True, text=True, check=True)
    return r.stdout.strip()


title = tex_inline(title)
abstract = tex_inline(abstract)
keywords = [tex_inline(k) for k in keywords]
run(["pandoc", "body.md", "-f", "markdown+tex_math_dollars", "-t", "latex", "-o", "body.tex", "--top-level-division=section"])
body_tex = (WORK / "body.tex").read_text()
# unnumbered headings are already numbered in the text ("1. Introduction") -> use starred sections
body_tex = re.sub(r"\\(section|subsection)\{", r"\\\1*{", body_tex)
body_tex = re.sub(r"\\label\{[^}]*\}", "", body_tex)
body_tex = body_tex.replace(r"\includegraphics[width=1\textwidth,height=\textheight]", r"\noindent\includegraphics[width=0.95\textwidth]")
# keep figure title, image and note together; keep table titles with their tables
body_tex = re.sub(r"(\\textbf\{Figure \d\..*?\})\n\n(\\noindent\\includegraphics\[[^]]*\]\{[^}]*\})\n\n(\\emph\{Note\}:.*?)\n\n",
                  lambda m: "\\begin{figure}[p]\n\\noindent " + m.group(1) + "\\par\\medskip\n" + m.group(2) + "\\par\\medskip\n{\\small " + m.group(3) + "}\n\\end{figure}\n\n",
                  body_tex, flags=re.S)
body_tex = re.sub(r"\n(\\textbf\{Table (?:\d+|A\.\d)\.)", r"\n\\needspace{14\\baselineskip}\1", body_tex)
# table column widths (formatter_agent rule): widen the label column, share the rest
def colspec(m):
    inner = m.group(0)[len(r"\begin{longtable}[]{@{}"):-len("@{}}")]
    n = inner.count("p{") or len(re.findall(r"[lrc]", inner))
    first = 0.34 if n <= 4 else (0.26 if n <= 6 else 0.20)
    rest = (1 - first) / (n - 1)
    gaps = 2 * (n - 1)
    cols = [first] + [rest] * (n - 1)
    spec = "".join(r">{\raggedright\arraybackslash}p{(\linewidth - %d\tabcolsep) * \real{%.4f}}" % (gaps, c) for c in cols)
    return r"\begin{longtable}[]{@{}" + spec + "@{}}"
body_tex = re.sub(r"\\begin\{longtable\}\[\]\{@\{\}.*?@\{\}\}", colspec, body_tex, flags=re.S)
# break long source paths and URLs (xurl)
body_tex = re.sub(r"(?<![{/])\b(output/[A-Za-z0-9_\\/.]+?\.csv)", lambda m: r"\url{" + m.group(1).replace("\\_", "_") + "}", body_tex)
body_tex = re.sub(r"(?<![{])(https?://[^\s}]+[^\s}.,;)])", lambda m: r"\url{" + m.group(1).replace("\\_", "_").replace("\\%", "%").replace("\\&", "&").replace("\\#", "#") + "}", body_tex)
tex = r"""\documentclass[review,authoryear,12pt]{elsarticle}
\usepackage{fontspec}
\setmainfont{TeX Gyre Termes}
\usepackage{amsmath,amssymb,graphicx,booktabs,longtable,array,calc,etoolbox}
\usepackage[hidelinks]{hyperref}
\usepackage{xurl}
\usepackage{setspace}
\usepackage{needspace}
\providecommand{\tightlist}{\setlength{\itemsep}{0pt}\setlength{\parskip}{0pt}}
\providecommand{\pandocbounded}[1]{#1}
\AtBeginEnvironment{longtable}{\footnotesize\setstretch{1.0}}
\setlength{\LTleft}{0pt}
\setlength{\emergencystretch}{3em}
\journal{Finance Research Open}
\begin{document}
\begin{frontmatter}
\title{""" + title + r"""}
\begin{abstract}
""" + abstract + r"""
\end{abstract}
\begin{keyword}
""" + r" \sep ".join(keywords) + r"""
\JEL """ + r" \sep ".join(jel) + r"""
\end{keyword}
\end{frontmatter}
\doublespacing
""" + body_tex + r"""
\end{document}
"""
tex = tex.replace("%", r"\%").replace(r"\\%", r"\%") if False else tex
(WORK / "manuscript_anonymized.tex").write_text(tex)
for _ in range(2):
    run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", "manuscript_anonymized.tex"])
log = (WORK / "manuscript_anonymized.log").read_text(errors="ignore")
overfull = re.findall(r"Overfull \\hbox \((\d+\.\d+)pt too wide\)", log)
print("overfull boxes > 10pt:", [x for x in overfull if float(x) > 10])

# ---------------------------------------------------------------- copy deliverables
for f in ("manuscript_anonymized.docx", "manuscript_anonymized.tex", "manuscript_anonymized.pdf", "title_page.docx",
          "highlights.docx", "cover_letter.docx", "declaration_of_competing_interest.docx", "manuscript_with_author_details.docx", "manuscript_anonymized.md"):
    shutil.copy(WORK / f, OUT / f)
print("done")
